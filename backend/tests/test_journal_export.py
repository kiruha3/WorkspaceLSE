from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from django.test import Client
from django.utils import timezone
from openpyxl import load_workbook

from apps.documents.models import Document
from apps.expertises.models import (
    Expertise, ExpertAssignment, HistoricalParticipant, HistoricalWork,
    ResultRevision, WorkCaseDetails, WorkClosure,
)
from apps.imports.models import ImportBatch, ImportRow
from apps.imports.source import parse_journal
from apps.intake.models import Material, Receipt
from apps.issuance.models import IssuePackage, MaterialIssue, PackageCheck, PackageRevision
from apps.references.models import Department, IssueMethod, Specialty

pytestmark = pytest.mark.django_db
URL = '/api/v1/expertises/export.xlsx'


@pytest.fixture
def registry(employee_factory, grant):
    actor = employee_factory(full_name='Эксперт А.А.')
    grant(actor, 'expertise.view')
    grant(actor, 'expertise.export')
    department = Department.objects.create(code='4', name='Отдел 4')
    specialty = Specialty.objects.create(code='18.1', name='Специальность')
    client = Client()
    client.force_login(actor.user)

    def create(number=151, **values):
        material = Material.objects.create(created_by=actor, title=f'Материал {number}')
        receipt = Receipt.objects.create(material=material, created_by=actor, received_at=date(2026, 2, 12))
        return Expertise.objects.create(
            material=material, receipt=receipt, registered_by=actor.user,
            registration_number=f'{number}/4-1-26', sequence=number, year=2026,
            category_code='1', department=department, specialty=specialty,
            appointment_nature='primary', **values,
        )
    return client, actor, department, specialty, create


def workbook(response):
    assert response.status_code == 200, response.content[:500]
    assert response['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'filename="expertises-' in response['Content-Disposition']
    return load_workbook(BytesIO(response.content))


def historical(work, actor, **values):
    batch = ImportBatch.objects.create(
        source_name='journal.csv', source_kind='csv', date_system='1900', source_sha256='a'*64,
        sheet_name='Лист1', parser_version='journal-2026-v1', uploaded_by=actor.user,
        row_count=1, summary={},
    )
    source = ImportRow.objects.create(batch=batch, source_row=4, raw={
        'A': '0042', 'B': 'STALE', 'D': 'Завершена', 'E': 'Старый орган',
        'J': 'старое дело', 'L': 'Старый эксперт', 'P': '01.01.2000',
        'Q': '99', 'V': '999', 'Y': '=HYPERLINK("https://invalid")',
        'Z': '999', 'AA': 'Устаревший результат', 'AB': 'свободная отметка',
    }, formulas={}, extra={}, facts={}, diagnostics=[])
    work.origin_kind = 'historical'
    work.save(update_fields=['origin_kind'])
    return HistoricalWork.objects.create(expertise=work, actor=actor.user, source_row=source, reason='Перенос', **values)


def test_export_retains_original_28_columns_three_row_header_and_is_reimportable(registry):
    client, actor, _, _, create = registry
    work = create(state='completed', historical_number='001/4-1-25')
    historical(work, actor, completed_on=date(2026, 3, 6), result_kind='expert_conclusion',
               questions=2, categorical=2, probable=0, unresolved=0, amount=Decimal('0.00'))
    book = workbook(client.get(URL))
    sheet = book.worksheets[0]
    assert sheet.title == 'Лист1'
    assert sheet.max_column == 28 and sheet.max_row == 4
    assert [sheet.cell(1, c).value for c in range(1, 29)] == [
        None, '№', 'Дата', 'состояние', 'Заказчик', None, None, None,
        'Документ основание', 'Номер дела, Ф.И.О. лиц, привлекаемых к ответственности, сторон по делу',
        'Материалы', None, 'N спец.', 'категория   сложности', 'результат', 'Дата окончания',
        'вопросов', 'выводы', None, None, None, 'Цена', 'Комментарий', 'отметки', None,
        'Срок производства', 'Результат', None,
    ]
    assert [sheet.cell(2, c).value for c in (5, 6, 7, 8, 18, 19, 20)] == [
        'тип органа', 'Орган', 'Лицо', 'контакт', 'категорич.', 'вероят.', 'не удалось',
    ]
    assert [sheet.cell(3, c).value for c in (12, 20, 21)] == ['Эксперт', 'кол-во', 'причина']
    assert {'E1:H1', 'R1:U1', 'T2:U2', 'B1:B3'} <= {str(r) for r in sheet.merged_cells.ranges}
    assert sheet.freeze_panes == 'E4'
    assert sheet.auto_filter.ref == 'A3:AB4'
    assert sheet['B4'].value == '151/4-1-26 (001/4-1-25)' and sheet['B4'].data_type == 's'
    assert sheet['M4'].value == '18.1' and sheet['M4'].data_type == 's'
    assert sheet['C4'].value.date() == date(2026, 2, 12)
    assert sheet['P4'].value.date() == date(2026, 3, 6)
    assert [sheet[f'{c}4'].value for c in ('Q','R','S','T','V')] == [2,2,0,0,0]
    parsed = parse_journal(client.get(URL).content, 'export.xlsx')
    assert parsed.rows[0].normalization_complete
    assert parsed.rows[0].facts['historical_number'] == '001/4-1-25'
    assert parsed.rows[0].facts['amount'] == '0.00'
    book.close()


def test_export_uses_current_records_and_only_falls_back_for_unmapped_cells(registry, employee_factory):
    client, actor, _, _, create = registry
    work = create(state='completed')
    history = historical(work, actor, completed_on=date(2026, 3, 6), result_kind='sndz',
                         questions=99, amount=None, notes='Текущий комментарий', delivery='Текущая отметка')
    WorkCaseDetails.objects.create(expertise=work, created_by=actor.user, authority_name='Новый орган',
                                   case_number='000123', case_description='Новое дело', contact='001234', inventory='Новые материалы')
    old = employee_factory(full_name='Бывший эксперт')
    HistoricalParticipant.objects.create(history=history, employee=old)
    ExpertAssignment.objects.create(expertise=work, employee=actor, actor=actor.user,
                                    started_at=timezone.now(), reason='Назначен')
    result = ResultRevision.objects.create(expertise=work, actor=actor.user, number=1, reason='Уточнено',
                                           kind='expert_conclusion', completed_on=date(2026, 4, 2),
                                           questions=2, categorical=2, probable=0, unresolved=None)
    document = Document.objects.create(kind='expert_conclusion', title='Заключение', source_kind='paper_original', created_by=actor)
    WorkClosure.objects.create(expertise=work, actor=actor.user, number=1, kind='completed',
                               ended_on=date(2026, 4, 2), result=result, document=document, reason='Завершена')
    ResultRevision.objects.create(expertise=work, actor=actor.user, number=2, reason='Непринятый черновик',
                                  kind='sndz', questions=999)
    book = workbook(client.get(URL)); sheet = book.active
    assert sheet['F4'].value == 'Новый орган' and sheet['E4'].value is None
    assert sheet['J4'].value == '000123\nНовое дело'
    assert sheet['H4'].value == '001234' and sheet['H4'].data_type == 's'
    assert sheet['K4'].value == 'Новые материалы' and sheet['L4'].value == actor.full_name
    assert sheet['P4'].value.date() == date(2026, 4, 2)
    assert [sheet[f'{c}4'].value for c in ('Q','R','S','T','V','W','X')] == [2,2,0,None,None,'Текущий комментарий','Текущая отметка']
    assert sheet['A4'].value == '0042'
    assert sheet['Y4'].value == '=HYPERLINK("https://invalid")' and sheet['Y4'].data_type == 's'
    assert sheet['AB4'].value == 'свободная отметка'
    assert sheet['Z4'].value == 50
    assert sheet['AA4'].value == 'Заключение эксперта 151/4-1-26 от 02.04.2026'
    assert not any(cell.data_type == 'f' for row in sheet for cell in row)
    book.close()


def test_export_keeps_large_price_exact_and_uses_corrected_issue_date(registry):
    client, actor, _, _, create = registry
    work = create(state='completed')
    historical(work, actor, amount=Decimal('9999999999999999.99'),
               notes='Выдано 01.01.2000', delivery='Материалы дела возвращены 01.01.2000',
               issued_on=date(2026, 4, 2))
    book = workbook(client.get(URL)); sheet = book.active
    assert sheet['V4'].value == '9999999999999999.99'
    assert sheet['V4'].data_type == 's'
    assert sheet['W4'].value == 'Выдано 02.04.2026'
    assert sheet['X4'].value == 'Выдано 02.04.2026'
    book.close()


def test_export_shares_calendar_and_employee_filters_with_registry(registry):
    client, actor, _, _, create = registry
    first = create(state='completed')
    history = historical(first, actor, completed_on=date(2026, 2, 12))
    HistoricalParticipant.objects.create(history=history, employee=actor)
    second = create(152, state='completed')
    historical(second, actor, completed_on=date(2026, 2, 13))
    filters = {'employee_id':str(actor.pk), 'calendar_days_min':0, 'calendar_days_max':1}
    assert client.get('/api/v1/expertises/', filters).json()['count'] == 1
    book = workbook(client.get(URL, filters))
    assert book.active.max_row == 4 and book.active['B4'].value == '151/4-1-26'
    book.close()


@pytest.mark.parametrize('description', ['x' * 32768, 'Недопустимый\x0bсимвол', 'Недопустимый\ufffeсимвол'], ids=['too-long','control-character','xml-noncharacter'])
def test_export_rejects_unrepresentable_cell_with_work_and_column_instead_of_losing_text(registry, description):
    client, actor, _, _, create = registry
    work = create()
    WorkCaseDetails.objects.create(expertise=work, created_by=actor.user, case_description=description)
    response = client.get(URL)
    assert response.status_code == 400
    error = response.json()
    assert '151/4-1-26' in str(error) and 'J' in str(error)
    assert response['Content-Type'].startswith('application/json')


def test_export_preserves_excel_cell_text_limit_and_legal_whitespace(registry):
    client, actor, _, _, create = registry
    work = create()
    description = '\t\n\r\r\n' + 'x' * 32762
    WorkCaseDetails.objects.create(expertise=work, created_by=actor.user, case_description=description)
    book = workbook(client.get(URL))
    assert book.active['J4'].value == description
    book.close()


def test_export_only_reveals_live_issue_details_when_all_material_works_allow_issue_view(registry, grant):
    client, actor, department, specialty, create = registry
    work = create(state='completed')
    historical(work, actor, notes='Исторический комментарий', completed_on=date(2026, 3, 6))
    package = IssuePackage.objects.create(material=work.material, created_by=actor.user)
    revision = PackageRevision.objects.create(package=package, number=1, actor=actor.user, reason='Комплект')
    checked = PackageCheck.objects.create(package=package, revision=revision, actor=actor.user, reason='Проверка')
    method = IssueMethod.objects.create(code='post', name='СЕКРЕТ способ')
    MaterialIssue.objects.create(package=package, package_check=checked, issued_on=date(2026, 4, 2),
                                 method=method, recipient_name='СЕКРЕТ получатель', recipient_basis='СЕКРЕТ основание',
                                 tracking='СЕКРЕТ трек', comment='СЕКРЕТ комментарий', actor=actor.user, reason='Выдано')
    # The date is already a visible expertise fact; delivery details have their own ACL.
    book = workbook(client.get(URL))
    assert book.active['W4'].value == 'Исторический комментарий'
    assert book.active['X4'].value == 'Выдано 02.04.2026'
    book.close()
    grant(actor, 'issue.view', 'departments', departments=[department])
    book = workbook(client.get(URL))
    assert book.active['W4'].value == 'Исторический комментарий\nСЕКРЕТ способ\nСЕКРЕТ получатель\nСЕКРЕТ основание\nСЕКРЕТ трек\nСЕКРЕТ комментарий'
    book.close()
    other = Department.objects.create(code='3', name='Другой отдел')
    Expertise.objects.create(material=work.material, receipt=work.receipt, registered_by=actor.user,
                             registration_number='152/3-1-26', sequence=152, year=2026, category_code='1',
                             department=other, specialty=specialty, appointment_nature='primary')
    # Filtering out a sibling does not make the shared package's protected facts visible.
    book = workbook(client.get(URL, {'search':'151/4-1-26'}))
    assert book.active.max_row == 4
    assert book.active['W4'].value == 'Исторический комментарий'
    assert book.active['X4'].value == 'Выдано 02.04.2026'
    book.close()


def test_export_does_not_resurrect_historical_experts_after_assignments_end(registry):
    client, actor, _, _, create = registry
    work = create(); history = historical(work, actor, executor_raw='Старый эксперт')
    HistoricalParticipant.objects.create(history=history, employee=actor)
    ExpertAssignment.objects.create(expertise=work, employee=actor, actor=actor.user,
                                    started_at=timezone.now(), ended_at=timezone.now(), reason='Окончено')
    book = workbook(client.get(URL))
    assert book.active['L4'].value is None
    assert book.active['P4'].value is None
    assert book.active['AA4'].value is None
    book.close()


def test_export_includes_all_filtered_rows_ignoring_pagination(registry):
    client, _, department, specialty, create = registry
    for number in range(151, 178): create(number, state='suspended')
    create(178, state='completed')
    response = client.get(URL, {'state':'suspended', 'department_id':str(department.pk),
                               'specialty_id':str(specialty.pk), 'received_from':'2026-02-12',
                               'received_to':'2026-02-12', 'search':'Материал', 'page':2, 'page_size':1})
    book = workbook(response)
    assert book.active.max_row == 30
    assert book.active['B4'].value == '177/4-1-26'
    assert book.active['B30'].value == '151/4-1-26'
    book.close()


@pytest.mark.parametrize('filters', [
    {'department_id':'invalid'}, {'state':'unknown'}, {'received_from':'bad'},
    {'calendar_days_min':-1}, {'calendar_days_min':10,'calendar_days_max':0},
    {'received_from':'2026-03-01','received_to':'2026-02-01'},
])
def test_export_validates_registry_filters(registry, filters):
    assert registry[0].get(URL, filters).status_code == 400


@pytest.mark.parametrize('actions', [(), ('expertise.view',), ('expertise.export',)])
def test_export_requires_both_view_and_existing_export_permission(registry, employee_factory, grant, actions):
    _, _, _, _, create = registry
    create()
    outsider = employee_factory()
    for action in actions: grant(outsider, action)
    client = Client(); client.force_login(outsider.user)
    assert client.get(URL).status_code == 403


def test_export_intersects_view_and_export_department_scopes(registry, employee_factory, grant):
    _, _, department, _, create = registry
    create(); other = Department.objects.create(code='3', name='Другой отдел')
    second = create(152); second.department = other; second.save(update_fields=['department'])
    viewer = employee_factory()
    grant(viewer, 'expertise.view', 'departments', departments=[department])
    grant(viewer, 'expertise.export')
    client = Client(); client.force_login(viewer.user)
    book = workbook(client.get(URL)); assert book.active.max_row == 4
    assert book.active['B4'].value == '151/4-1-26'; book.close()
    viewer.role_assignments.all().delete()
    grant(viewer, 'expertise.view')
    grant(viewer, 'expertise.export', 'departments', departments=[other])
    book = workbook(client.get(URL)); assert book.active.max_row == 4
    assert book.active['B4'].value == '152/4-1-26'; book.close()


def test_own_export_includes_historical_participants_only_until_assignments_exist(registry, employee_factory, grant):
    _, actor, _, _, create = registry
    work = create(); history = historical(work, actor)
    viewer = employee_factory()
    HistoricalParticipant.objects.create(history=history, employee=viewer)
    grant(viewer, 'expertise.view', 'own'); grant(viewer, 'expertise.export', 'own')
    client = Client(); client.force_login(viewer.user)
    book = workbook(client.get(URL)); assert book.active.max_row == 4; book.close()
    ExpertAssignment.objects.create(expertise=work, employee=actor, actor=actor.user,
                                    started_at=timezone.now(), ended_at=timezone.now(), reason='Окончено')
    book = workbook(client.get(URL)); assert book.active.max_row == 3; book.close()
