import pytest
from django.test import Client
from rest_framework.exceptions import ValidationError
from apps.common.ordering import ordered_rows

def test_natural_stable_sort_and_blanks_in_both_directions():
    rows=[{'value':'10/4-1-26','id':'a'},{'value':'2/4-1-26','id':'b'},{'value':None,'id':'c'},{'value':'2/4-1-26','id':'d'}]
    fields={'number':lambda row:row['value']}
    assert [row['id'] for row in ordered_rows(rows,'number',fields)]==['b','d','a','c']
    assert [row['id'] for row in ordered_rows(rows,'-number',fields)]==['a','b','d','c']
    assert ordered_rows(rows,'',fields)==rows
    assert [row['id'] for row in rows]==['a','b','c','d']
    with pytest.raises(ValidationError):ordered_rows(rows,'private_field',fields)

@pytest.mark.django_db
def test_employee_sorting_precedes_pagination_and_retains_filter_and_scope(employee_factory,grant):
    from apps.references.models import Department
    from apps.personnel.models import Employee
    department=Department.objects.create(code='3',name='Отдел 3')
    actor=employee_factory(department=department,full_name='Сотрудник 99')
    grant(actor,'employees.view','departments',departments=[department])
    for number in range(30,0,-1):
        employee=employee_factory(department=department,full_name=f'Сотрудник {number}')
    foreign=employee_factory(full_name='Сотрудник 0')
    client=Client();client.force_login(actor.user)
    first=client.get('/api/v1/employees/',{'ordering':'full_name'}).json()
    second=client.get('/api/v1/employees/',{'ordering':'full_name','page':2}).json()
    assert first['count']==31
    assert [item['full_name'] for item in first['results']][:3]==['Сотрудник 1','Сотрудник 2','Сотрудник 3']
    assert second['results'][0]['full_name']=='Сотрудник 26'
    assert str(foreign.id) not in [item['id'] for item in first['results']+second['results']]
    descending=client.get('/api/v1/employees/',{'ordering':'-full_name','search':'Сотрудник 2'}).json()
    assert [item['full_name'] for item in descending['results']]==[f'Сотрудник {i}' for i in range(29,19,-1)]+['Сотрудник 2']
    assert client.get('/api/v1/employees/',{'ordering':'user__password'}).status_code==400
    assert Employee.objects.count()==32

@pytest.mark.django_db
def test_expertise_numbers_and_export_share_natural_order_and_reset(employee_factory,grant):
    from datetime import date
    from io import BytesIO
    from openpyxl import load_workbook
    from apps.intake.models import Material,Receipt
    from apps.references.models import Department,Specialty
    from apps.expertises.models import Expertise
    actor=employee_factory();grant(actor,'expertise.view');grant(actor,'expertise.export')
    department=Department.objects.create(code='4',name='Отдел 4')
    specialty=Specialty.objects.create(code='18.1',name='Стоимость')
    for number in (10,2,1):
        material=Material.objects.create(created_by=actor,title=f'Материал {number}')
        receipt=Receipt.objects.create(material=material,created_by=actor,received_at=date(2026,1,number))
        Expertise.objects.create(material=material,receipt=receipt,registered_by=actor.user,registration_number=f'{number}/4-1-26',sequence=number,year=2026,category_code='1',department=department,specialty=specialty,appointment_nature='primary')
    client=Client();client.force_login(actor.user)
    expected=['1/4-1-26','2/4-1-26','10/4-1-26']
    for field in ('registration_number','received_at'):
        assert [r['registration_number'] for r in client.get('/api/v1/expertises/',{'ordering':field}).json()['results']]==expected
        assert [r['registration_number'] for r in client.get('/api/v1/expertises/',{'ordering':'-'+field}).json()['results']]==list(reversed(expected))
    assert [r['registration_number'] for r in client.get('/api/v1/expertises/').json()['results']]==list(reversed(expected))
    for field in ('department','executors','composition','state'):
        assert client.get('/api/v1/expertises/',{'ordering':field}).status_code==200
    response=client.get('/api/v1/expertises/export.xlsx',{'ordering':'registration_number'})
    assert response.status_code==200
    workbook=load_workbook(BytesIO(response.content),data_only=True)
    sheet=workbook.worksheets[0]
    values=[row[1].value for row in sheet.iter_rows() if len(row)>1 and row[1].value in expected]
    assert values==expected
    assert client.get('/api/v1/expertises/export.xlsx',{'ordering':'registration_number','search':'Нет такой работы'}).status_code==200
    assert client.get('/api/v1/expertises/',{'ordering':'material__created_by__user__password'}).status_code==400
