"""Bounded, read-only loading of the first journal worksheet."""
import csv
import hashlib
import re
from collections import Counter
from datetime import datetime,date
from io import BytesIO,StringIO
from itertools import zip_longest
from pathlib import PureWindowsPath
from zipfile import ZipFile,BadZipFile
from xml.etree import ElementTree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple
from .types import SourceRow,ParsedJournal
from .normalization import normalize_rows

MAX_FILE_BYTES=20*1024*1024
MAX_EXPANDED_BYTES=64*1024*1024
MAX_ROWS=10003
MAX_COLUMNS=128
PARSER_VERSION="journal-2026-v1"
HEADERS={"B":"№","C":"Дата","D":"состояние","K":"Материалы","M":"N спец.","N":"категория сложности","P":"Дата окончания","V":"Цена"}


class JournalSourceError(ValueError):
    pass


def cell_text(value):
    if value is None or value=="":return None
    if isinstance(value,datetime):return value.date().isoformat() if value.time()==datetime.min.time() else value.isoformat()
    if isinstance(value,date):return value.isoformat()
    return str(value)


def source_rows(records):
    rows=[]
    for index,(values,formulas) in enumerate(records,1):
        if index>MAX_ROWS:raise JournalSourceError("В источнике больше допустимых строк.")
        # Ignore style-only trailing cells, but never discard nonempty columns.
        if len(values)>MAX_COLUMNS and any(value is not None and str(value)!="" for value in values[MAX_COLUMNS:]):raise JournalSourceError("В источнике больше допустимых колонок.")
        if index==1:
            header={get_column_letter(i+1):value for i,value in enumerate(values)}
            clean=lambda value:re.sub(r"\s+","",str(value or "")).casefold()
            if any(clean(header.get(column))!=clean(label) for column,label in HEADERS.items()):raise JournalSourceError("Не распознаны заголовки первого листа журнала.")
        if index<=3:continue
        raw={get_column_letter(i+1):cell_text(values[i]) if i<len(values) else None for i in range(28)}
        # A and computed Z/AA/AB alone are template placeholders, not records.
        if not any(str(raw.get(column) or "").strip() or column in formulas for column in (get_column_letter(i+1) for i in range(1,24))):continue
        extra={get_column_letter(i+1):cell_text(value) for i,value in enumerate(values) if i>=28 and value is not None and str(value)!=""}
        rows.append(SourceRow(index,raw,formulas,extra))
    if not rows:raise JournalSourceError("В первом листе нет строк журнала.")
    return rows


def read_csv(content):
    try:
        decoded=content.decode("utf-8-sig")
        if "\x00" in decoded:raise JournalSourceError("CSV должен быть UTF-8 без нулевых байтов.")
        first=decoded.splitlines()[0] if decoded.splitlines() else ""
        delimiter=";" if first.count(";")>first.count(",") else ","
        def records():
            for values in csv.reader(StringIO(decoded,newline=""),delimiter=delimiter,strict=True):
                if len(values)>MAX_COLUMNS:raise JournalSourceError("В CSV больше допустимых колонок.")
                yield values,{}
        return source_rows(records()),"Лист1","1900"
    except (UnicodeError,csv.Error) as error:
        raise JournalSourceError("Не удалось прочитать UTF-8 CSV: проверьте кодировку и кавычки.") from error


def validate_zip(content):
    try:
        with ZipFile(BytesIO(content)) as archive:
            parts=archive.infolist()
            if len(parts)>1000 or sum(item.file_size for item in parts)>MAX_EXPANDED_BYTES:raise JournalSourceError("Раскрытый XLSX превышает допустимый размер.")
            for part in parts:
                if part.flag_bits&1:raise JournalSourceError("Зашифрованный XLSX не поддерживается.")
                if part.filename.casefold().endswith((".xml",".rels")):
                    xml=archive.read(part).replace(b"\x00",b"").upper()
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:raise JournalSourceError("XML с DTD/entity не поддерживается.")
    except (BadZipFile,KeyError,RuntimeError,OSError) as error:
        raise JournalSourceError("Файл не является читаемым XLSX.") from error


def sheet_date_serials(content,sheet):
    """Validate physical coordinates and retain numeric date values before coercion.

    openpyxl 3.1 read-only cells have already mapped styled serial 60 to Feb 28.
    This narrow XML pass preserves C/P values (including formula caches); it is
    not a second workbook reader. The pinned adapter uses its first worksheet's
    resolved package path so renamed/reordered worksheet parts still work.
    """
    dates={};row_number=0;column_number=0
    with ZipFile(BytesIO(content)) as archive,archive.open(sheet._worksheet_path) as stream:
        for event,element in ElementTree.iterparse(stream,events=("start","end")):
            tag=element.tag.rsplit("}",1)[-1]
            if event=="start" and tag=="row":
                row_number=int(element.get("r",row_number+1));column_number=0
                if not 1<=row_number<=MAX_ROWS:raise JournalSourceError("В источнике больше допустимых строк.")
            elif event=="end" and tag=="c":
                address=element.get("r")
                actual_row,column_number=coordinate_to_tuple(address) if address else (row_number,column_number+1)
                if actual_row!=row_number or not 1<=column_number<=MAX_COLUMNS:raise JournalSourceError("Координаты ячеек выходят за допустимые границы.")
                if column_number in (3,16) and element.get("t","n")=="n":
                    value=element.find("{*}v")
                    if value is not None and value.text is not None:dates[f"{get_column_letter(column_number)}{row_number}"]=value.text
                element.clear()
            elif event=="end" and tag=="row":element.clear()
    return dates


def read_xlsx(content):
    validate_zip(content)
    values_book=formula_book=None
    try:
        values_book=load_workbook(BytesIO(content),read_only=True,data_only=True,keep_links=False)
        formula_book=load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
        if not values_book.worksheets:raise JournalSourceError("В XLSX нет первого листа.")
        sheet=values_book.worksheets[0];formulas_sheet=formula_book.worksheets[0]
        dates=sheet_date_serials(content,sheet)
        # Producer-supplied dimensions may understate the data. Physical bounds
        # above are checked before iteration allocates or skips any cells.
        sheet.reset_dimensions();formulas_sheet.reset_dimensions()
        def records():
            for cells,source_cells in zip_longest(sheet.iter_rows(),formulas_sheet.iter_rows(),fillvalue=()):
                formulas={get_column_letter(i+1):str(cell.value) for i,cell in enumerate(source_cells) if cell.data_type=="f"}
                yield [dates.get(getattr(cell,"coordinate",None),cell.value) for cell in cells],formulas
        return source_rows(records()),sheet.title,"1904" if values_book.epoch.year==1904 else "1900"
    except JournalSourceError:raise
    except Exception as error:
        raise JournalSourceError("Не удалось прочитать структуру XLSX; проверьте исходный файл.") from error
    finally:
        if formula_book is not None:formula_book.close()
        if values_book is not None:values_book.close()


def parse_journal(content:bytes,filename:str)->ParsedJournal:
    if len(content)>MAX_FILE_BYTES:raise JournalSourceError("Размер источника превышает 20 MiB.")
    name=PureWindowsPath(filename).name
    kind=PureWindowsPath(name).suffix.casefold().lstrip(".")
    if kind not in ("csv","xlsx"):raise JournalSourceError("Поддерживаются только XLSX и CSV.")
    source,sheet,epoch=read_csv(content) if kind=="csv" else read_xlsx(content)
    rows=normalize_rows(source,date_system=epoch)
    summary={"row_count":len(rows),"error_rows":sum(not row.normalization_complete for row in rows),"warning_rows":sum(any(item.severity=="warning" for item in row.diagnostics) for row in rows),
        "diagnostic_counts":dict(Counter(item.code for row in rows for item in row.diagnostics)),"state_counts":dict(Counter(row.facts["state"] or "unknown" for row in rows))}
    return ParsedJournal(name,kind,epoch,hashlib.sha256(content).hexdigest(),sheet,PARSER_VERSION,rows,summary)
