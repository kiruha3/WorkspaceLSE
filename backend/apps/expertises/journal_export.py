"""A database snapshot in the original journal's first-worksheet layout.

Only A (the source ordinal) and the unlabeled Y/AB cells lack normalized
database fields. All other cells are built from current domain records.
No source formulas or stale raw values are copied into typed domain cells.
"""
from io import BytesIO
import os
import re
from zipfile import ZipFile

from django.utils import timezone
from openpyxl import LXML, Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError

from .history import historical_record
from .results import latest_closure, latest_result

MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
HEADERS = [
    None, '№', 'Дата', 'состояние', 'Заказчик', None, None, None,
    'Документ основание', 'Номер дела, Ф.И.О. лиц, привлекаемых к ответственности, сторон по делу',
    'Материалы', None, 'N спец.', 'категория   сложности', 'результат', 'Дата окончания',
    'вопросов', 'выводы', None, None, None, 'Цена', 'Комментарий', 'отметки', None,
    'Срок производства', 'Результат', None,
]
WIDTHS = [6.88,23,15.13,23.25,23.88,43.38,33.25,9.75,12.38,62,110,18.13,10.75,22.88,
          22.25,16.88,9.88,10.63,7.38,10.88,17.13,13.25,26.63,12.88,49.63,19.88,66.13,49.63]
MERGES = ['N1:N3','Z1:Z3','X1:X3','F2:F3','I1:I3','H2:H3','D1:D3','R1:U1','E2:E3',
          'AA1:AA3','J1:J3','P1:P3','A1:A3','V1:V3','K1:K3','T2:U2','E1:H1','R2:R3',
          'M1:M3','Q1:Q3','B1:B3','C1:C3','O1:O3','W1:W3','G2:G3','S2:S3']
STATES = {'completed':'Завершена', 'in_progress':'В работе', 'suspended':'Приостановлена',
          'without_execution':'Возврат без исполнения'}
RESULTS = {'expert_conclusion':'Заключение эксперта', 'research_act':'Акт экспертного исследования',
           'sndz':'СНДЗ'}
INVALID_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]')


def export_queryset(queryset):
    return queryset.select_related(
        'receipt', 'material', 'specialty', 'historical__source_row',
        'case_details__authority_type', 'material__issue_package__issuance__method',
    ).prefetch_related(
        'assignments__employee', 'historical__participants__employee',
        'historical__unresolved_allocations__reason', 'result_revisions__unresolved_allocations__reason',
        'closures__result__unresolved_allocations__reason',
    )


def joined(*values):
    return '\n'.join(str(value) for value in values if value is not None and str(value) != '') or None


def allocations_text(result):
    return joined(*(f'{item.reason.name}: {item.count}' + (f' — {item.comment}' if item.comment else '')
                    for item in result.unresolved_allocations.all()))


def issue_marker(value):
    return bool(re.fullmatch(r'(?:выдано|материалы\s+дела\s+возвращены)\s+\d{2}\.\d{2}\.\d{4}\.?', value or '', re.I))


def exact_amount(value):
    # Excel numeric cells have only 15 significant decimal digits.
    return format(value, '.2f') if value is not None and len(value.as_tuple().digits) > 15 else value


def journal_values(work, ordinal, as_of, *, can_view_issue=False):
    history = historical_record(work)
    raw = history.source_row.raw if history else {}
    details = getattr(work, 'case_details', None)
    closed = work.state in ('completed', 'without_execution')
    closure = latest_closure(work) if closed else None
    result = closure.result if closure else latest_result(work)
    # A present current result/closure is authoritative even when its fields are unknown.
    result_facts = result if result is not None else None if closure else history
    end = (closure.ended_on if closure else
           history.completed_on if history and work.state == 'completed' else
           history.return_decided_on if history and work.state == 'without_execution' else None)
    if result is not None and work.state == 'completed' and closure is None:
        end = result.completed_on
    result_kind = getattr(result_facts, 'kind', None) if result is not None else getattr(result_facts, 'result_kind', None)
    result_label = RESULTS.get(result_kind)
    registration = work.registration_number + (f' ({work.historical_number})' if work.historical_number else '')
    assignments = list(work.assignments.all())
    if assignments:
        executors = joined(*(item.employee.full_name for item in assignments if item.ended_at is None))
    elif history:
        participants = list(history.participants.all())
        executors = joined(*(item.employee.full_name for item in participants)) if participants else history.executor_raw
    else:
        executors = None
    case = details if details is not None else history
    def case_value(field):
        return getattr(case, field, None)
    authority_type = (details.authority_type.name if details.authority_type_id else None) if details else case_value('authority_type')
    unresolved_reason = allocations_text(result_facts) if result_facts else None
    if result is None and closure is None and history:
        unresolved_reason = joined(unresolved_reason or history.unresolved_reason_raw, history.unresolved_note)
    package = getattr(work.material, 'issue_package', None)
    issue = getattr(package, 'issuance', None) if package else None
    issued_on = issue.issued_on if issue else history.issued_on if history else None
    delivery = history.delivery if history else None
    notes = history.notes if history else None
    if issue_marker(notes):
        notes = None
    if issue_marker(delivery):
        delivery = None
    if issued_on:
        # Keep one standalone factual marker that the same parser can read back.
        marker = f'Выдано {issued_on:%d.%m.%Y}'
        if issue:
            if can_view_issue:
                notes = joined(notes, issue.method.name, issue.recipient_name, issue.recipient_basis, issue.tracking, issue.comment)
            delivery = marker
        elif history and issue_marker(history.notes):
            notes = marker
            if issue_marker(history.delivery):
                delivery = marker
        elif not delivery:
            delivery = marker
        elif not notes:
            notes = marker
        else:
            notes = joined(notes, delivery)
            delivery = marker
    result_reference = f'{result_label} {registration} от {end:%d.%m.%Y}' if result_label and end else None
    through = end if closed else as_of
    return [
        raw.get('A') if history else ordinal,
        registration, work.receipt.received_at, STATES[work.state], authority_type,
        case_value('authority_name') if case else work.receipt.sender_name,
        case_value('official'), case_value('contact'), case_value('basis_text'),
        joined(getattr(details, 'case_number', None), case_value('case_description')),
        case_value('inventory'), executors, work.specialty.code if work.specialty_id else None,
        getattr(result_facts, 'complexity_level', None), result_label, end,
        *(getattr(result_facts, field, None) for field in ('questions','categorical','probable','unresolved')),
        unresolved_reason, exact_amount(history.amount) if history else None, notes, delivery, raw.get('Y'),
        max((through - work.receipt.received_at).days + 1, 0) if through else None,
        result_reference, raw.get('AB'),
    ]


def build_journal(queryset, *, issue_material_ids=()):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Лист1'
    sheet.append(HEADERS)
    for address, value in {'E2':'тип органа','F2':'Орган','G2':'Лицо','H2':'контакт',
                           'R2':'категорич.','S2':'вероят.','T2':'не удалось',
                           'L3':'Эксперт','T3':'кол-во','U3':'причина'}.items():
        sheet[address] = value
    for merge in MERGES:
        sheet.merge_cells(merge)
    side = Side(style='thin', color='A0A0A0')
    border = Border(left=side, right=side, top=side, bottom=side)
    for row in sheet.iter_rows(min_row=1, max_row=3, max_col=28):
        for cell in row:
            cell.font = Font(name='Arial', size=10, bold=True)
            cell.fill = PatternFill('solid', fgColor='C9C9C9')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
    for row in range(1, 4):
        sheet.row_dimensions[row].height = 30
    as_of = timezone.localdate()
    issue_material_ids = frozenset(issue_material_ids)
    for ordinal, work in enumerate(export_queryset(queryset).iterator(chunk_size=250), 1):
        values = journal_values(work, ordinal, as_of, can_view_issue=work.material_id in issue_material_ids)
        for column, value in enumerate(values, 1):
            if not isinstance(value, str):
                continue
            location = f'Экспертиза {work.registration_number}, колонка {get_column_letter(column)}'
            if len(value) > 32767:
                raise ValidationError(f'{location}: текст превышает предел Excel — 32767 символов. Данные не обрезаны.')
            if INVALID_XML.search(value):
                raise ValidationError(f'{location}: текст содержит символ, недопустимый для Excel. Уточните значение в карточке.')
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                # XLSX string cells do not execute =/+/-/@ payloads; no visible apostrophe needed.
                cell.data_type = 's'
                cell.number_format = '@'
            cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        for column in ('C', 'P'):
            sheet[f'{column}{sheet.max_row}'].number_format = 'DD.MM.YYYY'
        sheet[f'V{sheet.max_row}'].number_format = '#,##0.00'
    for column, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = 'E4'
    sheet.auto_filter.ref = f'A3:AB{max(sheet.max_row, 3)}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:3'
    output = BytesIO()
    book.save(output)
    book.close()
    # ElementTree writes raw CR in text; XML readers normalize it to LF.
    # A character reference preserves the actual DB text in Excel and reimport.
    escaped = BytesIO()
    with ZipFile(output) as source, ZipFile(escaped, 'w') as target:
        for part in source.infolist():
            content = source.read(part)
            if part.filename.startswith('xl/worksheets/') and part.filename.endswith('.xml'):
                if not LXML and os.linesep == '\r\n':
                    # et_xmlfile's Windows text stream expanded original LF to CRLF.
                    content = content.replace(b'\r\n', b'\n')
                content = content.replace(b'\r', b'&#13;')
            target.writestr(part, content)
    return escaped.getvalue()
